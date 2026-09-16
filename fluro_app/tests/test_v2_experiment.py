# -*- coding: utf-8 -*-
"""v2 功能改进测试集：单卡片槽比色检测规格。

问题 1：实验信息配置（实验参数持久化 + 校验，默认 T/Bg 双 ROI）。
后续问题 2~4 在此文件逐步追加。
"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import pytest  # noqa: E402
import numpy as np  # noqa: E402

from app import create_app  # noqa: E402
from app import database as db  # noqa: E402


@pytest.fixture()
def client(tmp_path):
    app = create_app({'TESTING': True, 'DATA_DIR': str(tmp_path)})
    return app.test_client()


# ============ 问题 1：实验信息配置 ============

def test_p1_default_experiment(client):
    ex = client.get('/api/experiment').get_json()
    assert ex['carrier'] == 'single_card'
    assert ex['roi_count'] == 2
    assert ex['roi_names'] == ['T', 'Bg']
    assert ex['has_blank_ref'] is True
    assert 'analyte' in ex and 'color_trend' in ex


def test_p1_save_and_load(client):
    payload = {
        'analyte': '某农药',
        'color_trend': '浓度越高 T 区越深',
        'roi_count': 2,
        'roi_names': ['T', 'Bg'],
        'has_color_card': False,
        'has_blank_ref': True,
        'known_concs': '0,1,2,5,10,20,50',
        'replicates': '3',
        'image_count': '21',
        'note': '单卡片槽',
    }
    r = client.put('/api/experiment', json=payload)
    assert r.status_code == 200
    saved = client.get('/api/experiment').get_json()
    assert saved['analyte'] == '某农药'
    assert saved['known_concs'] == '0,1,2,5,10,20,50'
    assert saved['replicates'] == '3'


def test_p1_validate_roi_names(client):
    assert client.put('/api/experiment', json={'roi_names': []}).status_code == 400
    assert client.put('/api/experiment', json={'roi_names': ['T', 'T']}).status_code == 400
    assert client.put('/api/experiment', json={'roi_names': ['']}).status_code == 400


def test_p1_validate_roi_count(client):
    assert client.put('/api/experiment', json={'roi_count': 0}).status_code == 400
    assert client.put('/api/experiment', json={'roi_count': 11}).status_code == 400
    assert client.put('/api/experiment', json={'roi_count': 'x'}).status_code == 400


def test_p1_validate_concs(client):
    assert client.put('/api/experiment', json={'known_concs': '0,a,3'}).status_code == 400
    assert client.put('/api/experiment', json={'known_concs': '0,-1,3'}).status_code == 400
    r = client.put('/api/experiment', json={'known_concs': '0, 1.5, 2'})
    assert r.status_code == 200
    assert r.get_json()['experiment']['known_concs'] == '0,1.5,2'


def test_p1_validate_replicates(client):
    assert client.put('/api/experiment', json={'replicates': 'x'}).status_code == 400
    assert client.put('/api/experiment', json={'replicates': '0'}).status_code == 400
    assert client.put('/api/experiment', json={'replicates': '3'}).status_code == 200


def test_p1_persists_in_settings(client):
    client.put('/api/experiment', json={'analyte': '持久化测试', 'roi_names': ['T', 'Bg']})
    # 直接查 settings 表：重启后仍可读（数据持久化，不丢）
    dpath = Path(client.application.config['DATA_DIR']) / 'fluro.db'
    row = db.query_one(dpath, "SELECT value FROM settings WHERE key='experiment'")
    assert row is not None
    assert '持久化测试' in row['value']


def test_p1_roi_names_saved_for_modeling(client):
    """ROI 命名（T/Bg）是建模自动识别的基础，保存后必须完整读回。"""
    client.put('/api/experiment', json={'roi_names': ['T', 'Bg']})
    ex = client.get('/api/experiment').get_json()
    assert ex['roi_names'] == ['T', 'Bg']


# ============ 问题 2-A：多命名 ROI 引擎（T/Bg） ============

ROIS_T_BG = [
    {'name': 'T', 'role': 'sample', 'x': 0.335, 'y': 0.28, 'w': 0.33, 'h': 0.44, 'bg_subtract': 0},
    {'name': 'Bg', 'role': 'background', 'x': 0.10, 'y': 0.10, 'w': 0.20, 'h': 0.30, 'bg_subtract': 0},
]


def _upload_img(client, conc=50, kind='detection'):
    import io
    import cv2 as _cv2
    from tools import make_test_images as _mti
    img = _mti.make_test_image(conc)
    ok, buf = _cv2.imencode('.png', img[:, :, ::-1])
    data = {'files': [(io.BytesIO(buf.tobytes()), 't.png')], 'kind': kind}
    if kind == 'calibration':
        data['known_conc'] = str(conc)
    return client.post('/api/images/upload', data=data, content_type='multipart/form-data').get_json()['images'][0]['id']


def test_p2a_save_and_get_multi_rois(client):
    iid = _upload_img(client)
    r = client.post(f'/api/images/{iid}/rois', json={'rois': ROIS_T_BG})
    assert r.status_code == 200
    got = client.get(f'/api/images/{iid}/rois').get_json()['rois']
    assert len(got) == 2
    names = sorted(x['name'] for x in got)
    assert names == ['Bg', 'T']
    t = next(x for x in got if x['name'] == 'T')
    assert t['role'] == 'sample'
    bg = next(x for x in got if x['name'] == 'Bg')
    assert bg['role'] == 'background'


def test_p2a_multi_roi_validation(client):
    iid = _upload_img(client)
    assert client.post(f'/api/images/{iid}/rois', json={'rois': []}).status_code == 400
    dup = [dict(ROIS_T_BG[0]), dict(ROIS_T_BG[0])]
    dup[0]['name'] = 'T'
    assert client.post(f'/api/images/{iid}/rois', json={'rois': dup}).status_code == 400
    bad = [dict(ROIS_T_BG[0], x=1.5)]
    assert client.post(f'/api/images/{iid}/rois', json={'rois': bad}).status_code == 400
    missing = [{'name': 'T'}]
    assert client.post(f'/api/images/{iid}/rois', json={'rois': missing}).status_code == 400


def test_p2a_main_roi_synced_to_legacy(client):
    """多 ROI 保存后，主检测区（T）同步到旧 roi 表，兼容原有特征/检测链路。"""
    iid = _upload_img(client)
    client.post(f'/api/images/{iid}/rois', json={'rois': ROIS_T_BG})
    legacy = client.get(f'/api/pipeline/{iid}/roi').get_json()['roi']
    assert legacy is not None
    assert abs(legacy['x'] - ROIS_T_BG[0]['x']) < 1e-6


def test_p2a_features_use_main_roi(client):
    iid = _upload_img(client, conc=50)
    client.post('/api/pipeline/run', json={'image_ids': [iid]})
    client.post(f'/api/images/{iid}/rois', json={'rois': ROIS_T_BG})
    f = client.get(f'/api/pipeline/{iid}/features').get_json()['features']
    assert f['mean_r'] is not None and f['hue'] is not None


def test_p2a_template_multi_roi_apply(client):
    iid = _upload_img(client)
    tpl = {
        'name': '卡片模板',
        'template_json': {
            'T': {'x': 0.335, 'y': 0.28, 'w': 0.33, 'h': 0.44, 'role': 'sample'},
            'Bg': {'x': 0.10, 'y': 0.10, 'w': 0.20, 'h': 0.30, 'role': 'background'},
        },
    }
    tid = client.post('/api/templates', json=tpl).get_json()['id']
    r = client.post(f'/api/images/{iid}/rois/apply_template', json={'template_id': tid})
    assert r.status_code == 200
    got = client.get(f'/api/images/{iid}/rois').get_json()['rois']
    assert len(got) == 2 and sorted(x['name'] for x in got) == ['Bg', 'T']
    assert all(x['source'] == 'template' for x in got)


def test_p2a_apply_legacy_single_template_as_t(client):
    iid = _upload_img(client)
    tid = client.post('/api/templates', json={'name': '单ROI', 'x': 0.3, 'y': 0.3, 'w': 0.4, 'h': 0.4}).get_json()['id']
    r = client.post(f'/api/images/{iid}/rois/apply_template', json={'template_id': tid})
    assert r.status_code == 200
    got = client.get(f'/api/images/{iid}/rois').get_json()['rois']
    assert len(got) == 1 and got[0]['name'] == 'T'


def test_p2a_overlay_all_rois(client):
    iid = _upload_img(client)
    client.post(f'/api/images/{iid}/rois', json={'rois': ROIS_T_BG})
    r = client.get(f'/api/images/{iid}/overlay')
    assert r.status_code == 200 and r.headers['Content-Type'].startswith('image/png')


# ============ 问题 2-B：扩展特征提取（RGB/SD、Lab、OD、ΔE、比值） ============

def test_p2b_extract_roi_features_full_fields(client):
    from app.image_processing import extract_roi_features
    import numpy as _np
    from tools import make_test_images as _mti
    img = _mti.make_test_image(50)
    f = extract_roi_features(img, (0.335, 0.28, 0.33, 0.44))
    for key in ('mean_r', 'mean_g', 'mean_b', 'median_r', 'median_g', 'median_b',
                'std_r', 'std_g', 'std_b', 'hue', 'saturation', 'value',
                'lab_l', 'lab_a', 'lab_b', 'gray', 'od_r', 'od_g', 'od_b',
                'ratio_gr', 'ratio_gb', 'ratio_rb'):
        assert f[key] is not None, f'{key} 缺失'
    # OD 单调性：亮像素 OD 小，暗像素 OD 大
    bright = _np.full((20, 20, 3), 240, dtype=_np.uint8)
    dark = _np.full((20, 20, 3), 30, dtype=_np.uint8)
    assert extract_roi_features(bright, (0, 0, 1, 1))['od_r'] < extract_roi_features(dark, (0, 0, 1, 1))['od_r']


def test_p2b_compute_multi_roi_features(client):
    iid = _upload_img(client, conc=100)  # T 区红色、Bg 区取背景附近
    client.post(f'/api/images/{iid}/rois', json={'rois': ROIS_T_BG})
    r = client.post(f'/api/images/{iid}/features', json={})
    assert r.status_code == 200
    d = r.get_json()
    assert 'T' in d['rois'] and 'Bg' in d['rois']
    t = d['rois']['T']
    assert t['lab_l'] is not None and t['od_g'] is not None
    comb = d['combined']
    assert comb.get('T_R_over_Bg_R') is not None
    assert comb.get('T_R_minus_Bg_R') is not None
    assert comb.get('deltaE_T_vs_Bg') is not None
    # 主 ROI 兼容旧特征表
    legacy = client.get(f'/api/pipeline/{iid}/features').get_json()['features']
    assert legacy['hue'] is not None


def test_p2b_roi_features_persisted(client):
    iid = _upload_img(client, conc=50)
    client.post(f'/api/images/{iid}/rois', json={'rois': ROIS_T_BG})
    client.post(f'/api/images/{iid}/features', json={})
    got = client.get(f'/api/images/{iid}/features').get_json()
    assert set(got['rois'].keys()) == {'T', 'Bg'}
    assert 'deltaE_T_vs_Bg' in got['combined']
    assert got['combined']['T_G_over_Bg_G'] > 0


def test_p2b_derive_delta_e_symmetric(client):
    """色差 ΔE 应为对称且非负（Lab 空间欧氏距离）。"""
    from app.image_processing import derive_combined_features
    f1 = {'mean_r': 200, 'mean_g': 50, 'mean_b': 50, 'od_r': 0.1, 'od_g': 0.7, 'od_b': 0.7,
          'lab_l': 60, 'lab_a': 70, 'lab_b': 40}
    f2 = {'mean_r': 90, 'mean_g': 90, 'mean_b': 90, 'od_r': 0.5, 'od_g': 0.5, 'od_b': 0.5,
          'lab_l': 70, 'lab_a': 0, 'lab_b': 0}
    c1 = derive_combined_features({'T': f1, 'Bg': f2})
    c2 = derive_combined_features({'T': f2, 'Bg': f1})
    assert c1['deltaE_T_vs_Bg'] == c2['deltaE_T_vs_Bg']
    assert c1['deltaE_T_vs_Bg'] > 0


def test_p2b_features_require_roi(client):
    iid = _upload_img(client)
    r = client.post(f'/api/images/{iid}/features', json={})
    assert r.status_code == 400
    assert 'ROI' in r.get_json()['error']


# ============ 问题 2-C：白平衡 / 白参考校正 ============

def test_p2c_white_balance_correct_unit():
    from app.image_processing import white_balance_correct
    import numpy as _np
    img = _np.full((10, 10, 3), [100, 100, 100], dtype=_np.uint8)
    # 白参考偏暗（如 Bg 暗背景）→ 增益 >1，图像提亮
    brightened = white_balance_correct(img, [50, 50, 50])
    assert int(brightened[0, 0, 0]) > 100
    # 白参考接近 255 → 基本不变
    unchanged = white_balance_correct(img, [250, 250, 250])
    assert abs(int(unchanged[0, 0, 0]) - 100) <= 2


def test_p2c_preprocess_with_white_balance(client):
    iid = _upload_img(client, conc=50)
    client.post(f'/api/images/{iid}/rois', json={'rois': ROIS_T_BG})
    # 不开白平衡
    r1 = client.post(f'/api/pipeline/{iid}/preprocess',
                     json={'filter': 'gaussian', 'kernel': 5, 'use_dark': False, 'use_flat': False,
                           'use_wb': False})
    assert r1.status_code == 200
    img1 = _decode_resp(client.get(f'/api/images/{iid}/processed'))
    # 开白平衡（Bg 作白参考）
    r2 = client.post(f'/api/pipeline/{iid}/preprocess',
                     json={'filter': 'gaussian', 'kernel': 5, 'use_dark': False, 'use_flat': False,
                           'use_wb': True, 'wb_roi_name': 'Bg'})
    assert r2.status_code == 200
    assert r2.get_json()['params']['use_wb'] is True
    assert r2.get_json()['params']['wb_roi_name'] == 'Bg'
    img2 = _decode_resp(client.get(f'/api/images/{iid}/processed'))
    diff = np.abs(img1.astype(int) - img2.astype(int))
    assert diff.mean() > 5, '白平衡校正未改变处理结果'


def test_p2c_white_balance_requires_roi(client):
    iid = _upload_img(client, conc=50)
    r = client.post(f'/api/pipeline/{iid}/preprocess',
                    json={'filter': 'gaussian', 'kernel': 5, 'use_wb': True})
    assert r.status_code == 400
    assert '白参考' in r.get_json()['error']


def _decode_resp(resp):
    import cv2 as _cv2
    import numpy as _np
    arr = _np.frombuffer(resp.data, dtype=_np.uint8)
    return _cv2.imdecode(arr, _cv2.IMREAD_COLOR)[:, :, ::-1]


# ============ 问题 2-D：图片导入增强（文件夹/格式/文件信息） ============

def test_p2d_upload_keeps_filename(client):
    """上传时保留原始文件名，缩略图列表可读（问题2 界面1 的文件信息要求）。"""
    import io
    import cv2 as _cv2
    from tools import make_test_images as _mti
    img = _mti.make_test_image(20)
    ok, buf = _cv2.imencode('.png', img[:, :, ::-1])
    data = {'files': [(io.BytesIO(buf.tobytes()), 'sample_T1_20ng.png')], 'kind': 'detection',
            'batch': 'B-2D'}
    res = client.post('/api/images/upload', data=data, content_type='multipart/form-data').get_json()
    assert res['images'][0]['filename'] == 'sample_T1_20ng.png'
    listing = client.get('/api/images').get_json()['images']
    assert any(im['filename'] == 'sample_T1_20ng.png' for im in listing)


def test_p2d_supported_formats(client):
    """jpg/jpeg/png/bmp/tif/tiff/webp 均可导入；非图片文件被拒且不影响其余文件。"""
    import io
    import cv2 as _cv2
    from tools import make_test_images as _mti
    img = _mti.make_test_image(10)
    files = []
    for ext in ('.jpg', '.png', '.bmp', '.tif', '.webp'):
        ok, buf = _cv2.imencode(ext, img[:, :, ::-1])
        assert ok, f'{ext} 编码失败'
        files.append((io.BytesIO(buf.tobytes()), 'img' + ext))
    files.append((io.BytesIO(b'not an image'), 'bad.txt'))
    res = client.post('/api/images/upload', data={'files': files, 'kind': 'detection'},
                      content_type='multipart/form-data').get_json()
    assert len(res['images']) == 5
    assert res['errors'] and 'bad.txt' in res['errors'][0]


# ============ 问题 2-E：结果表可编辑 + CSV/Excel 导出 + 界面1→2 数据 ============

def test_p2e_edit_image_meta(client):
    """结果表可编辑：批次/浓度/重复编号/备注（PATCH 持久化）。"""
    iid = _upload_img(client, conc=50)
    r = client.patch(f'/api/images/{iid}',
                     json={'batch': 'B-2E', 'known_conc': 42.5, 'replicate': 3, 'note': '三次重复'})
    assert r.status_code == 200
    row = client.get('/api/images').get_json()['images'][0]
    assert row['batch'] == 'B-2E'
    assert abs(row['known_conc'] - 42.5) < 1e-6
    assert row['replicate'] == 3
    assert row['note'] == '三次重复'
    # 校验：非法浓度 / 重复编号
    assert client.patch(f'/api/images/{iid}', json={'known_conc': 'abc'}).status_code == 400
    assert client.patch(f'/api/images/{iid}', json={'replicate': 0}).status_code == 400


def test_p2e_export_csv_spec_columns(client):
    """CSV 导出符合单卡片规格列：image_name/concentration/T_*/Bg_*/ΔE/比值/OD。"""
    import csv as _csv
    import io as _io
    iid = _upload_img(client, conc=20, kind='calibration')
    client.post(f'/api/images/{iid}/rois', json={'rois': ROIS_T_BG})
    client.post(f'/api/images/{iid}/features', json={})
    client.patch(f'/api/images/{iid}', json={'batch': 'EXP-1', 'replicate': 2, 'note': 'ok'})
    r = client.get('/api/export/features.csv')
    assert r.status_code == 200
    text = r.data.decode('utf-8-sig')
    reader = _csv.DictReader(_io.StringIO(text))
    cols = reader.fieldnames
    for c in ('image_name', 'image_path', 'batch', 'concentration', 'replicate',
              'T_R', 'T_G', 'T_B', 'T_H', 'T_S', 'T_V', 'T1_L', 'T1_a', 'T1_b',
              'Bg_R', 'Bg_G', 'Bg_B', 'deltaE_T_vs_Bg',
              'T_R_over_Bg_R', 'T_G_over_Bg_G', 'T_B_over_Bg_B',
              'OD_T_R', 'OD_T_G', 'OD_T_B', 'note'):
        assert c in cols, f'缺少列 {c}'
    rows = list(reader)
    assert len(rows) == 1
    row = rows[0]
    assert row['image_name']  # filename 已保留
    assert row['concentration'] == '20.0' or abs(float(row['concentration']) - 20) < 1e-6
    assert row['batch'] == 'EXP-1' and row['replicate'] == '2'
    assert float(row['T_R']) > 0 and float(row['deltaE_T_vs_Bg']) > 0
    assert float(row['T_R_over_Bg_R']) > 0


def test_p2e_export_xlsx(client):
    """Excel 导出可被 openpyxl 打开且含规格列。"""
    import openpyxl
    import io as _io
    iid = _upload_img(client, conc=10, kind='calibration')
    client.post(f'/api/images/{iid}/rois', json={'rois': ROIS_T_BG})
    client.post(f'/api/images/{iid}/features', json={})
    r = client.get('/api/export/features.xlsx')
    assert r.status_code == 200
    wb = openpyxl.load_workbook(_io.BytesIO(r.data))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert 'T_R' in header and 'deltaE_T_vs_Bg' in header
    assert ws.max_row == 2  # 表头 + 一行数据


def test_p2e_export_empty_table_ok(client):
    """无特征数据时导出表头完整、零数据行（不报错）。"""
    r = client.get('/api/export/features.csv')
    assert r.status_code == 200
    text = r.data.decode('utf-8-sig')
    assert 'image_name' in text
    assert text.count('\n') == 1  # 只有表头
